"""
Document Builder - Creates DOCX directly from analysis structure.

Uses python-docx like test_table_generator.py instead of markdown->Pandoc.
"""
import json
from pathlib import Path
from typing import Dict, Any, Optional
from docx import Document
from docx.shared import Pt
from loguru import logger

# Add parent directory to path for imports
import sys
sys.path.append(str(Path(__file__).parent.parent))

from ..builders.docx_builder import build_table_from_data


def build_document_from_analysis(
    analysis_text: str,
    sources_data: Dict[str, Dict[str, Any]],
    task_id: str = "unknown",
    task_folder: Optional[str] = None
) -> str:
    """
    Build DOCX document directly from analysis structure.
    Same logic as test_table_generator.py create_full_document_from_json
    """
    logger.info(f"[{task_id}] Building document from analysis structure...")
    
    # Parse JSON directly (like test_table_generator)
    try:
        content = analysis_text.strip()
        if content.startswith('```json'):
            content = content[7:]
        elif content.startswith('```'):
            content = content[3:]
        if content.endswith('```'):
            content = content[:-3]
        content = content.strip()
        
        full_json = json.loads(content)
        logger.info(f"[{task_id}] Loaded JSON with {len(full_json)} items")
    except json.JSONDecodeError as e:
        logger.error(f"[{task_id}] Failed to parse JSON: {e}")
        raise ValueError(f"Invalid JSON in analysis: {e}")
    
    # Create new document
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(14)
    
    # Process all items in order (exactly like test_table_generator)
    for item in full_json:
        item_type = item.get('type', 'text')
        title = item.get('title', '')
        level = item.get('level', 0)
        
        if item_type == 'text':
            # Process text section
            description = item.get('description', {})
            
            # Build content from description
            content_parts = []
            if isinstance(description, dict):
                general = description.get('general', '')
                specific = description.get('specific', '')
                if specific.startswith('## Наприклад:'):
                    specific = specific[13:].strip()
                if general:
                    content_parts.append(general)
                if specific:
                    content_parts.append(specific)
            
            content_text = '\n\n'.join(content_parts) if content_parts else f"[{title}]"
            
            # Add heading with appropriate style
            if level == 0:
                heading_style = 'Title'
            elif level == 1:
                heading_style = 'Heading 1'
            elif level == 2:
                heading_style = 'Heading 2'
            elif level == 3:
                heading_style = 'Heading 3'
            elif level == 4:
                heading_style = 'Heading 4'
            else:
                heading_style = 'Heading 5'
            
            doc.add_paragraph(title, style=heading_style)
            
            # Add content paragraphs
            paragraphs = content_text.split('\n\n')
            for para_text in paragraphs:
                if para_text.strip():
                    doc.add_paragraph(para_text.strip())
            
            doc.add_paragraph()  # Empty line after section
            logger.info(f"[{task_id}] Added section: {title}")
            
        elif item_type == 'table':
            # Process table directly from JSON (like test_table_generator)
            columns = item.get('columns', [])
            if isinstance(columns, list):
                headers = columns
            else:
                headers = []
            
            # Add table title as heading
            if level == 0:
                heading_style = 'Title'
            elif level == 1:
                heading_style = 'Heading 1'
            elif level == 2:
                heading_style = 'Heading 2'
            elif level == 3:
                heading_style = 'Heading 3'
            else:
                heading_style = 'Heading 4'
            
            doc.add_paragraph(title, style=heading_style)
            
            # Create table with empty rows (user will fill data)
            if headers:
                build_table_from_data(
                    doc,
                    item.get('id', 0),
                    headers,
                    [],  # Empty rows
                    None
                )
                logger.info(f"[{task_id}] Added table: {title} with {len(headers)} columns")
            
            doc.add_paragraph()  # Empty line after table
    
    # Determine output path
    if task_folder:
        output_dir = Path(task_folder) / 'temp'
        output_dir.mkdir(exist_ok=True)
        output_path = output_dir / 'temp_result.docx'
    else:
        output_path = Path('/tmp') / f'document_{task_id}.docx'
    
    # Save document
    doc.save(str(output_path))
    logger.success(f"[{task_id}] Document saved: {output_path}")
    
    return str(output_path)
