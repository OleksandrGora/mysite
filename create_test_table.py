#!/usr/bin/env python3
"""
Create a test DOCX file with tables for testing table parsing
"""
import sys
import os
sys.path.append('/home/openclaw/mysite')

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

def create_test_docx():
    """Create a test DOCX file with tables."""
    doc = Document()
    
    # Add title
    title = doc.add_heading('Тестовий документ з таблицями', 0)
    
    # Add some text
    doc.add_paragraph('Це тестовий документ для перевірки парсингу таблиць.')
    
    # Add section heading
    doc.add_heading('Розділ 1: Органолептичні показники', level=1)
    
    # Create table 1
    table1 = doc.add_table(rows=6, cols=5)
    table1.style = 'Table Grid'
    
    # Add headers
    headers = ['Показник', 'Характеристика', 'Яєчного порошку', 'Білка яєчного сухого', 'Жовтка яєчного сухого']
    for i, header in enumerate(headers):
        cell = table1.cell(0, i)
        cell.text = header
    
    # Add data rows
    data_rows = [
        ['Зовнішній вигляд', '描述', 'Сипучий порошок', 'Сипучий порошок', 'Сипучий порошок'],
        ['Колір', '描述', 'Світло-жовтий', 'Світло-жовтий', 'Жовто-помаранчевий'],
        ['Запах', '描述', 'Властивий яєчному порошку', 'Властивий яєчному білку', 'Властивий яєчному жовтку'],
        ['Вологість, %', 'Не більше 7.0', '6.5', '6.8', '7.0'],
        ['Розчинність', 'Повна', 'Повна', 'Повна', 'Повна']
    ]
    
    for row_idx, row_data in enumerate(data_rows, start=1):
        for col_idx, cell_data in enumerate(row_data):
            table1.cell(row_idx, col_idx).text = cell_data
    
    # Add some text between tables
    doc.add_paragraph('\nТаблиця 1 показує основні органолептичні показники сухих яєчних продуктів.')
    
    # Add section heading for second table
    doc.add_heading('Розділ 2: Хімічний склад', level=1)
    
    # Create table 2
    table2 = doc.add_table(rows=4, cols=3)
    table2.style = 'Table Grid'
    
    # Add headers
    headers2 = ['Компонент', 'Яєчний порошок', 'Білок яєчний сухий']
    for i, header in enumerate(headers2):
        table2.cell(0, i).text = header
    
    # Add data
    data2 = [
        ['Протеїн, %', '45.0-48.0', '82.0-85.0'],
        ['Жир, %', '38.0-42.0', '1.5-3.0'],
        ['Вуглеводи, %', '3.0-5.0', '4.0-6.0']
    ]
    
    for row_idx, row_data in enumerate(data2, start=1):
        for col_idx, cell_data in enumerate(row_data):
            table2.cell(row_idx, col_idx).text = cell_data
    
    # Save the document
    output_path = '/home/openclaw/mysite/data/templates/test_table.docx'
    doc.save(output_path)
    print(f"Test DOCX file created: {output_path}")
    
    return output_path

if __name__ == '__main__':
    create_test_docx()
