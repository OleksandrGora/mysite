"""
txt_processor - Shared text extraction module for various file formats.
Provides a unified interface for extracting text from documents, spreadsheets,
presentations, markup files, and plain text files.

Supported formats:
  - .docx (Word)          - via python-docx
  - .pdf  (PDF)           - via PyMuPDF, with OCR fallback for scanned PDFs
  - .doc  (Word 97)       - via olefile
  - .txt  (plain text)    - with encoding fallback (utf-8 / cp1251)
  - .md   (markdown)      - plain text extraction
  - .html/.htm (HTML)     - via html.parser (built-in) with encoding fallback
  - .csv  (CSV)           - via csv module (built-in)
  - .json (JSON)          - formatted via json module
  - .pptx (PowerPoint)    - via python-pptx
  - .xlsx (Excel)         - via openpyxl

Usage:
    from txt_processor import extract_text, extract_text_from_bytes

    # From file path
    text = extract_text('/path/to/document.docx')

    # From bytes (e.g. uploaded file)
    text = extract_text_from_bytes(file_bytes, 'report.docx')
"""

import csv
import io
import json
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Dict, List, Optional

from loguru import logger


# ---------------------------------------------------------------------------
# HTML text extractor (uses only stdlib)
# ---------------------------------------------------------------------------

class _HTMLTextExtractor(HTMLParser):
    """HTML parser that strips tags and returns plain text."""

    def __init__(self):
        super().__init__()
        self._parts = []
        self._skip_tag = False

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style'):
            self._skip_tag = True

    def handle_endtag(self, tag):
        if tag in ('script', 'style'):
            self._skip_tag = False

    def handle_data(self, data):
        if not self._skip_tag:
            self._parts.append(data)

    def get_text(self):
        return ''.join(self._parts)


# ---------------------------------------------------------------------------
# Text cleaning
# ---------------------------------------------------------------------------

def clean_text(text: str) -> str:
    """Normalize line endings, collapse excess whitespace."""
    if not text:
        return ''
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    text = re.sub(r' +', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


# ---------------------------------------------------------------------------
# ZIP bomb protection for archive-based formats (docx, pptx, xlsx)
# ---------------------------------------------------------------------------

_ZIP_MAX_COMPRESSION_RATIO = 100
_ZIP_MAX_UNCOMPRESSED = 100 * 1024 * 1024  # 100 MB
_DOC_MAX_STREAM_SIZE = 100 * 1024 * 1024  # 100 MB

def _check_zip_bomb(file_path: str) -> None:
    """Raise ValueError if a ZIP-based file exhibits ZIP bomb characteristics."""
    try:
        if not zipfile.is_zipfile(file_path):
            return
    except Exception:
        return

    with zipfile.ZipFile(file_path) as zf:
        total_compressed, total_uncompressed = _sum_zip_sizes(zf)

    _verify_zip_sizes(total_compressed, total_uncompressed)


def _check_zip_bomb_bytes(data: bytes) -> None:
    """Raise ValueError if in-memory ZIP data exhibits ZIP bomb characteristics."""
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        total_compressed, total_uncompressed = _sum_zip_sizes(zf)

    _verify_zip_sizes(total_compressed, total_uncompressed)


def _sum_zip_sizes(zf: zipfile.ZipFile):
    """Return (total_compressed, total_uncompressed) for all entries in a ZipFile."""
    total_compressed = 0
    total_uncompressed = 0
    for info in zf.infolist():
        total_compressed += info.compress_size
        total_uncompressed += info.file_size
    return total_compressed, total_uncompressed


def _verify_zip_sizes(total_compressed: int, total_uncompressed: int) -> None:
    """Raise ValueError if compressed/uncompressed sizes exceed bomb thresholds."""
    if total_compressed > 0:
        ratio = total_uncompressed / total_compressed
        if ratio > _ZIP_MAX_COMPRESSION_RATIO:
            raise ValueError(
                f"Potential ZIP bomb: compression ratio {ratio:.0f}:1 exceeds "
                f"{_ZIP_MAX_COMPRESSION_RATIO}:1 limit"
            )
    if total_uncompressed > _ZIP_MAX_UNCOMPRESSED:
        raise ValueError(
            f"Potential ZIP bomb: uncompressed size {total_uncompressed:,} bytes "
            f"exceeds {_ZIP_MAX_UNCOMPRESSED:,} limit"
        )


# ---------------------------------------------------------------------------
# OCR fallback (isolated cross-blueprint import)
# ---------------------------------------------------------------------------

def _try_ocr_fallback(file_path: str) -> Optional[str]:
    """Attempt OCR on a scanned PDF via the doc_generator blueprint.

    The import is kept isolated in this function so that a missing or
    restructured blueprint does not prevent the rest of txt_processor from
    loading.
    """
    try:
        from blueprints.doc_generator.parsers import ocr
    except ImportError:
        logger.debug("OCR module not available for PDF fallback")
        return None
    try:
        text = ocr.parse_pdf_with_ocr(file_path)
        if text:
            return clean_text(text)
    except Exception as e:
        logger.warning(f"OCR fallback for PDF failed: {e}")
    return None


# ---------------------------------------------------------------------------
# Per-format extractors (from file path)
# ---------------------------------------------------------------------------

def _extract_docx(file_path: str) -> str:
    _check_zip_bomb(file_path)
    from docx import Document
    doc = Document(file_path)
    return _extract_docx_core(doc)


def _extract_docx_core(doc) -> str:
    """Shared core: extract text from an already-opened python-docx Document."""
    paragraphs = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            paragraphs.append(t)
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                rows.append(' | '.join(cells))
        if rows:
            paragraphs.append('\n'.join(rows))
    return clean_text('\n\n'.join(paragraphs))


def _extract_pdf(file_path: str) -> str:
    import fitz
    doc = fitz.open(file_path)
    try:
        pages = []
        total_chars = 0
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text()
            total_chars += len(text.strip())
            if text.strip():
                pages.append(text)

        avg_chars = total_chars / max(len(pages), 1)
        if avg_chars < 100:
            ocr_text = _try_ocr_fallback(file_path)
            if ocr_text:
                return ocr_text

        return clean_text('\n\n'.join(pages))
    finally:
        doc.close()


def _extract_pdf_core(doc) -> str:
    """Shared core: extract text from an already-opened PyMuPDF document (no OCR)."""
    pages = []
    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text()
        if text.strip():
            pages.append(text)
    return clean_text('\n\n'.join(pages))


def _extract_doc(file_path: str) -> str:
    import olefile
    ole = olefile.OleFileIO(file_path)
    try:
        stream = ole.openstream('WordDocument')
        _check_ole_stream_size(stream)
        data = stream.read()
    finally:
        ole.close()
    return _extract_doc_core(data)


def _check_ole_stream_size(stream) -> None:
    """Raise ValueError if OLE stream exceeds the maximum allowed size."""
    try:
        size = stream.size
    except AttributeError:
        return
    if size > _DOC_MAX_STREAM_SIZE:
        raise ValueError(
            f"OLE WordDocument stream too large: {size:,} bytes exceeds "
            f"{_DOC_MAX_STREAM_SIZE:,} limit"
        )


def _extract_doc_core(data: bytes) -> str:
    """Shared core: extract text from raw .doc OLE stream bytes."""
    raw = data.decode('utf-16-le', errors='replace')

    def _is_printable(c):
        return (c.isalpha() or c.isspace() or c.isdigit()
                or c in '.,;:!?\'"„”“’–—()-—±×÷°′″№«»'
                or 0x0400 <= ord(c) <= 0x04FF)

    regions = []
    start = None
    length = 0
    for i, c in enumerate(raw):
        if _is_printable(c):
            if start is None:
                start = i
                length = 1
            else:
                length += 1
        else:
            if length >= 20:
                regions.append((start, length, raw[start:start + length]))
            start = None
            length = 0
    if length >= 20:
        regions.append((start, length, raw[start:start + length]))

    if not regions:
        raise ValueError("No text found in .doc file")

    regions.sort(key=lambda r: r[1], reverse=True)
    lines = [l.strip() for l in regions[0][2].split('\n') if len(l.strip()) > 3]
    return clean_text('\n'.join(lines))


def _extract_txt(file_path: str) -> str:
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            return clean_text(f.read())
    except UnicodeDecodeError:
        pass
    try:
        with open(file_path, 'r', encoding='cp1251') as f:
            return clean_text(f.read())
    except UnicodeDecodeError:
        pass
    with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
        return clean_text(f.read())


def _extract_md(file_path: str) -> str:
    return _extract_txt(file_path)


def _extract_html(file_path: str, data: Optional[bytes] = None) -> str:
    if data is not None:
        raw = data
    else:
        with open(file_path, 'rb') as f:
            raw = f.read()
    content = _decode_text_bytes(raw)
    parser = _HTMLTextExtractor()
    parser.feed(content)
    return clean_text(parser.get_text())


def _extract_csv(file_path: str, data: Optional[bytes] = None) -> str:
    if data is not None:
        raw = data
    else:
        with open(file_path, 'rb') as f:
            raw = f.read()
    content = _decode_text_bytes(raw)
    rows = [' | '.join(row) for row in csv.reader(io.StringIO(content)) if any(c.strip() for c in row)]
    return clean_text('\n'.join(rows))


def _extract_json(file_path: str, data: Optional[bytes] = None) -> str:
    if data is not None:
        raw_text = _decode_text_bytes(data)
        try:
            parsed = json.loads(raw_text)
        except json.JSONDecodeError:
            logger.warning("JSON parsing failed, returning raw text")
            return clean_text(raw_text)
    else:
        with open(file_path, 'rb') as f:
            raw_bytes = f.read()
        raw_text = _decode_text_bytes(raw_bytes)
        try:
            parsed = json.loads(raw_text)
        except json.JSONDecodeError:
            logger.warning(f"JSON parsing failed for {file_path}, returning raw text")
            return clean_text(raw_text)
    return json.dumps(parsed, indent=2, ensure_ascii=False)


def _extract_pptx(file_path: str) -> str:
    _check_zip_bomb(file_path)
    from pptx import Presentation
    prs = Presentation(file_path)
    return _extract_pptx_core(prs)


def _extract_pptx_core(prs) -> str:
    """Shared core: extract text from an already-opened python-pptx Presentation."""
    paragraphs = []
    for slide_num, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append(t)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        texts.append(' | '.join(cells))
        if texts:
            paragraphs.append(f"[Slide {slide_num}]\n" + '\n'.join(texts))
    return clean_text('\n\n'.join(paragraphs))


def _extract_xlsx(file_path: str) -> str:
    _check_zip_bomb(file_path)
    from openpyxl import load_workbook
    wb = load_workbook(file_path, read_only=True, data_only=True)
    return _extract_xlsx_core(wb)


def _extract_xlsx_core(wb) -> str:
    """Shared core: extract text from an already-opened openpyxl Workbook."""
    sheets = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(' | '.join(cells))
        if rows:
            sheets.append(f"[Sheet: {sheet_name}]\n" + '\n'.join(rows))
    wb.close()
    return clean_text('\n\n'.join(sheets))


# ---------------------------------------------------------------------------
# Text decoding helper (encoding fallback chain)
# ---------------------------------------------------------------------------

def _decode_text_bytes(raw: bytes) -> str:
    """Decode bytes to str, trying common encodings with full fallback."""
    for enc in ('utf-8', 'cp1251', 'latin-1'):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return raw.decode('utf-8', errors='replace')


# ---------------------------------------------------------------------------
# Bytes-based extractors (no temp-file round-trip)
# ---------------------------------------------------------------------------

def _extract_docx_bytes(data: bytes) -> str:
    _check_zip_bomb_bytes(data)
    from docx import Document
    doc = Document(io.BytesIO(data))
    return _extract_docx_core(doc)


def _extract_pdf_bytes(data: bytes) -> str:
    import fitz
    doc = fitz.open(stream=data, filetype='pdf')
    try:
        return _extract_pdf_core(doc)
    finally:
        doc.close()


def _extract_doc_bytes(data: bytes) -> str:
    import olefile
    ole = olefile.OleFileIO(data=data)
    try:
        stream = ole.openstream('WordDocument')
        _check_ole_stream_size(stream)
        doc_data = stream.read()
    finally:
        ole.close()
    return _extract_doc_core(doc_data)


def _extract_txt_bytes(data: bytes) -> str:
    return clean_text(_decode_text_bytes(data))


def _extract_md_bytes(data: bytes) -> str:
    return clean_text(_decode_text_bytes(data))


def _extract_html_bytes(data: bytes) -> str:
    return _extract_html('', data=data)


def _extract_csv_bytes(data: bytes) -> str:
    return _extract_csv('', data=data)


def _extract_json_bytes(data: bytes) -> str:
    return _extract_json('', data=data)


def _extract_pptx_bytes(data: bytes) -> str:
    _check_zip_bomb_bytes(data)
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    return _extract_pptx_core(prs)


def _extract_xlsx_bytes(data: bytes) -> str:
    _check_zip_bomb_bytes(data)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    return _extract_xlsx_core(wb)


# ---------------------------------------------------------------------------
# Extension registry
# ---------------------------------------------------------------------------

EXTENSION_MAP: Dict[str, callable] = {
    '.docx': _extract_docx,
    '.pdf': _extract_pdf,
    '.doc': _extract_doc,
    '.txt': _extract_txt,
    '.md': _extract_md,
    '.html': _extract_html,
    '.htm': _extract_html,
    '.csv': _extract_csv,
    '.json': _extract_json,
    '.pptx': _extract_pptx,
    '.xlsx': _extract_xlsx,
}

BYTES_EXTENSION_MAP: Dict[str, callable] = {
    '.docx': _extract_docx_bytes,
    '.pdf': _extract_pdf_bytes,
    '.doc': _extract_doc_bytes,
    '.txt': _extract_txt_bytes,
    '.md': _extract_md_bytes,
    '.html': _extract_html_bytes,
    '.htm': _extract_html_bytes,
    '.csv': _extract_csv_bytes,
    '.json': _extract_json_bytes,
    '.pptx': _extract_pptx_bytes,
    '.xlsx': _extract_xlsx_bytes,
}

# Extensions that contain extractable text (for upload filtering)
SUPPORTED_EXTENSIONS: set = set(EXTENSION_MAP.keys())

# MIME type mapping for document types
MIME_MAP: Dict[str, str] = {
    '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    '.pdf': 'application/pdf',
    '.doc': 'application/msword',
    '.txt': 'text/plain',
    '.md': 'text/markdown',
    '.html': 'text/html',
    '.htm': 'text/html',
    '.csv': 'text/csv',
    '.json': 'application/json',
    '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
    '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
}

# Image extensions (not processed by txt_processor, but kept for reference)
IMAGE_EXTENSIONS: set = {'.jpg', '.jpeg', '.png', '.gif', '.webp'}


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def extract_text(file_path: str) -> str:
    """
    Extract text from a file based on its extension.

    Args:
        file_path: Path to the source file.

    Returns:
        Extracted and cleaned text content.

    Raises:
        ValueError: If the file type is not supported or parsing fails.
    """
    ext = Path(file_path).suffix.lower()
    if ext in IMAGE_EXTENSIONS:
        raise ValueError(f"Image files ({ext}) are not supported for text extraction")
    handler = EXTENSION_MAP.get(ext)
    if handler is None:
        raise ValueError(f"Unsupported file type: {ext}")
    logger.info(f"Extracting text from {file_path} (type: {ext})")
    return handler(file_path)


def extract_text_from_bytes(data: bytes, filename: str) -> str:
    """
    Extract text from file content in memory (e.g. uploaded file).

    Uses in-memory extraction for all supported formats, avoiding
    unnecessary disk round-trips.

    Args:
        data: Raw file content.
        filename: Original filename (used to determine extension).

    Returns:
        Extracted and cleaned text content.

    Raises:
        ValueError: If the file type is not supported or parsing fails.
    """
    ext = Path(filename).suffix.lower()
    handler = BYTES_EXTENSION_MAP.get(ext)
    if handler is None:
        raise ValueError(f"Unsupported file type: {ext}")
    logger.info(f"Extracting text from bytes: {filename} (type: {ext})")
    return handler(data)


def extract_texts_from_files(file_paths: List[str]) -> Dict[str, str]:
    """
    Extract text from multiple files.

    Returns:
        Dict mapping filename -> extracted text.
        On error the value is a short error message.
    """
    results = {}
    for path in file_paths:
        try:
            text = extract_text(path)
            results[Path(path).name] = text
        except Exception as e:
            logger.error(f"Failed to extract text from {path}: {e}")
            results[Path(path).name] = f"[Error: {e}]"
    return results


def is_supported(filename: str) -> bool:
    """Check if a filename has a supported text-extraction extension."""
    return Path(filename).suffix.lower() in EXTENSION_MAP


def is_image(filename: str) -> bool:
    """Check if a filename is a supported image type."""
    return Path(filename).suffix.lower() in IMAGE_EXTENSIONS
